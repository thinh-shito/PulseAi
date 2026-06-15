"""Document processing utilities for the MCP server."""

import base64
import io
import logging
from typing import Union

logger = logging.getLogger(__name__)


def process_file_content(file_bytes: bytes, filename: str) -> dict:
    """
    Process a file (docx or pdf) and return its text or image content.

    Args:
        file_bytes: Raw file bytes
        filename: Original filename to determine type

    Returns:
        dict with 'type' ('text' or 'images') and 'content'
    """
    ext = filename.lower().split(".")[-1] if "." in filename else ""

    if ext == "docx":
        return _process_docx(file_bytes)
    elif ext == "pdf":
        return _process_pdf(file_bytes)
    else:
        # Assume plain text
        try:
            text = file_bytes.decode("utf-8", errors="replace")
            return {"type": "text", "content": text}
        except Exception as e:
            logger.error(f"Failed to decode file as text: {e}")
            return {"type": "text", "content": ""}


def _process_docx(file_bytes: bytes) -> dict:
    """Extract text from a DOCX file."""
    try:
        import docx2txt
        from docx import Document

        # Try docx2txt first for better text extraction
        try:
            text = docx2txt.process(io.BytesIO(file_bytes))
            if text and text.strip():
                return {"type": "text", "content": text.strip()}
        except Exception:
            pass

        # Fallback to python-docx
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs)
        return {"type": "text", "content": text}
    except Exception as e:
        logger.error(f"Error processing DOCX: {e}")
        return {"type": "text", "content": ""}


def _process_pdf(file_bytes: bytes) -> dict:
    """Extract text from a PDF file. Falls back to image rendering if needed."""
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(file_bytes))
        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)

        text = "\n".join(text_parts).strip()
        if text:
            return {"type": "text", "content": text}

        # If no text extracted, render as images
        return _render_pdf_as_images(file_bytes)

    except Exception as e:
        logger.error(f"Error processing PDF text: {e}")
        return _render_pdf_as_images(file_bytes)


def _render_pdf_as_images(file_bytes: bytes) -> dict:
    """Render PDF pages as base64 JPEG images."""
    try:
        from pdf2image import convert_from_bytes

        images = convert_from_bytes(file_bytes, dpi=150)
        image_b64_list = []
        for img in images:
            buf = io.BytesIO()
            img.save(buf, format="JPEG", quality=85)
            b64 = base64.b64encode(buf.getvalue()).decode("utf-8")
            image_b64_list.append(b64)

        return {"type": "images", "content": image_b64_list}
    except Exception as e:
        logger.error(f"Error rendering PDF as images: {e}")
        return {"type": "images", "content": []}


def fill_docx_template(template_path: str, field_values: dict) -> bytes:
    """
    Fill a DOCX template by replacing {{placeholder}} patterns with values.

    Args:
        template_path: Path to the .docx template file
        field_values: Dict mapping placeholder names to values

    Returns:
        Bytes of the filled DOCX file
    """
    import re
    from docx import Document
    from docx.oxml import OxmlElement

    def sanitize(value: str) -> str:
        """Strip XML-incompatible control characters."""
        invalid_xml = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]")
        return invalid_xml.sub("", str(value))

    def replace_in_paragraph(para, replacements: dict):
        """Replace placeholders within a paragraph, handling split runs."""
        # Rebuild full paragraph text
        full_text = "".join(run.text for run in para.runs)
        new_text = full_text
        for key, value in replacements.items():
            placeholder = "{{" + key + "}}"
            if placeholder in new_text:
                new_text = new_text.replace(placeholder, sanitize(str(value) if value else ""))

        if new_text != full_text:
            # Assign all text to first run, clear others
            if para.runs:
                para.runs[0].text = new_text
                for run in para.runs[1:]:
                    run.text = ""

    # Normalize field values: strip quotes, replace blank values
    cleaned = {}
    for k, v in field_values.items():
        if isinstance(v, str):
            v = v.strip().strip('"').strip("'")
        if v in ("", "N/A", "n/a", "None", None):
            v = " "
        cleaned[k] = v

    doc = Document(template_path)

    # Replace in paragraphs
    for para in doc.paragraphs:
        replace_in_paragraph(para, cleaned)

    # Replace in tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    replace_in_paragraph(para, cleaned)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()