"""Tool: generate_handover_report — Generate a clinical handover report for shift transitions."""

import base64
import json
import logging
import os
from pathlib import Path

from openai import OpenAI

logger = logging.getLogger(__name__)

HANDOVER_SYSTEM_PROMPT = """You are a clinical documentation specialist. Generate a structured shift handover report based on the provided patient information.

The handover report must follow this exact JSON structure:
{
  "patient_name": "Full patient name",
  "patient_id": "MRN or patient ID",
  "ward": "Ward/unit name",
  "bed_number": "Bed number",
  "attending_physician": "Name of attending physician",
  "handover_date": "Date of handover (YYYY-MM-DD)",
  "handover_time": "Time of handover (HH:MM)",
  "outgoing_nurse": "Name of outgoing nurse/clinician",
  "incoming_nurse": "Name of incoming nurse/clinician (if provided)",
  "primary_diagnosis": "Primary diagnosis",
  "active_problems": ["List of active clinical problems"],
  "vitals_summary": "Summary of recent vital signs (BP, HR, RR, Temp, SpO2)",
  "current_medications": ["List of current medications with doses"],
  "recent_interventions": ["List of interventions performed this shift"],
  "pending_tasks": ["List of tasks that must be completed next shift"],
  "critical_alerts": ["Any critical alerts or concerns requiring immediate attention"],
  "pain_score": "Current pain score (0-10)",
  "diet_restrictions": "Current diet or fluid restrictions",
  "mobility_status": "Current mobility and fall risk status",
  "lines_and_drains": ["List of IV lines, drains, catheters"],
  "additional_notes": "Any other relevant clinical notes"
}

Extract as much information as possible from the provided clinical summary. Use empty strings or empty arrays for fields that cannot be determined. Return only valid JSON."""


def generate_handover_report(clinical_summary: str, patient_info: str = "") -> dict:
    """
    Generate a structured clinical handover report for shift transitions.

    Args:
        clinical_summary: Clinical summary text describing patient status this shift
        patient_info: Optional additional patient demographic/contextual information

    Returns:
        dict with structured handover report data and a generated DOCX file
    """
    client = OpenAI(api_key=os.environ.get("OPENAI_API_KEY"))
    model = os.environ.get("OPENAI_MODEL", "gpt-4o")

    combined_input = clinical_summary
    if patient_info:
        combined_input = f"## Patient Information\n{patient_info}\n\n## Clinical Summary\n{clinical_summary}"

    logger.info("Generating handover report via OpenAI")

    try:
        response = client.chat.completions.create(
            model=model,
            temperature=0.0,
            messages=[
                {"role": "system", "content": HANDOVER_SYSTEM_PROMPT},
                {"role": "user", "content": combined_input},
            ],
            response_format={"type": "json_object"},
        )

        result_text = response.choices[0].message.content
        report_data = json.loads(result_text)

        logger.info("Handover report generated successfully")

        # Generate a text-based DOCX document
        docx_b64, filename = _create_handover_docx(report_data)

        return {
            "report": report_data,
            "filename": filename,
            "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "content_base64": docx_b64,
            "message": f"Handover report generated successfully: {filename}",
        }

    except Exception as e:
        logger.error(f"Error generating handover report: {e}")
        raise


def _create_handover_docx(report_data: dict) -> tuple[str, str]:
    """Create a formatted DOCX handover report from structured data."""
    import io
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Title
    title = doc.add_heading("Clinical Handover Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Header info
    doc.add_paragraph(
        f"Patient: {report_data.get('patient_name', 'N/A')}  |  "
        f"MRN: {report_data.get('patient_id', 'N/A')}  |  "
        f"Ward: {report_data.get('ward', 'N/A')}  |  "
        f"Bed: {report_data.get('bed_number', 'N/A')}"
    )
    doc.add_paragraph(
        f"Date: {report_data.get('handover_date', 'N/A')}  |  "
        f"Time: {report_data.get('handover_time', 'N/A')}  |  "
        f"Attending: {report_data.get('attending_physician', 'N/A')}"
    )
    doc.add_paragraph(
        f"Outgoing: {report_data.get('outgoing_nurse', 'N/A')}  |  "
        f"Incoming: {report_data.get('incoming_nurse', 'N/A')}"
    )
    doc.add_paragraph()

    # Sections
    def add_section(heading: str, content):
        doc.add_heading(heading, level=2)
        if isinstance(content, list):
            if content:
                for item in content:
                    doc.add_paragraph(f"• {item}", style="List Bullet")
            else:
                doc.add_paragraph("None documented.")
        else:
            doc.add_paragraph(str(content) if content else "Not documented.")

    add_section("Primary Diagnosis", report_data.get("primary_diagnosis", ""))
    add_section("Active Problems", report_data.get("active_problems", []))
    add_section("Vitals Summary", report_data.get("vitals_summary", ""))
    add_section("Current Medications", report_data.get("current_medications", []))
    add_section("Recent Interventions", report_data.get("recent_interventions", []))
    add_section("Lines & Drains", report_data.get("lines_and_drains", []))

    # Critical alerts section (highlighted)
    alerts = report_data.get("critical_alerts", [])
    doc.add_heading("⚠️ Critical Alerts", level=2)
    if alerts:
        for alert in alerts:
            p = doc.add_paragraph()
            run = p.add_run(f"⚠ {alert}")
            run.bold = True
    else:
        doc.add_paragraph("No critical alerts.")

    add_section("Pending Tasks (Next Shift)", report_data.get("pending_tasks", []))
    add_section("Pain Score", report_data.get("pain_score", ""))
    add_section("Diet & Fluid Restrictions", report_data.get("diet_restrictions", ""))
    add_section("Mobility Status", report_data.get("mobility_status", ""))
    add_section("Additional Notes", report_data.get("additional_notes", ""))

    buf = io.BytesIO()
    doc.save(buf)
    doc_bytes = buf.getvalue()

    patient_name = report_data.get("patient_name", "patient").replace(" ", "_")
    date = report_data.get("handover_date", "").replace("-", "") or "undated"
    filename = f"handover_report_{patient_name}_{date}.docx"

    return base64.b64encode(doc_bytes).decode("utf-8"), filename