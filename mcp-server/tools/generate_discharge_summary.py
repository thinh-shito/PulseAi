"""Tool: generate_discharge_summary — Generate a clinical discharge summary document."""

import base64
import json
import logging
import os

from openai import OpenAI

logger = logging.getLogger(__name__)

DISCHARGE_SYSTEM_PROMPT = """You are a clinical documentation specialist. Generate a comprehensive discharge summary based on the provided patient information.

The discharge summary must follow this exact JSON structure:
{
  "patient_name": "Full patient name",
  "patient_id": "MRN or patient ID",
  "patient_dob": "Date of birth (YYYY-MM-DD)",
  "admission_date": "Date of admission (YYYY-MM-DD)",
  "discharge_date": "Date of discharge (YYYY-MM-DD)",
  "length_of_stay": "Number of days hospitalized",
  "attending_physician": "Name of attending physician",
  "ward": "Ward/unit where patient was admitted",
  "discharge_disposition": "Where patient is being discharged to (e.g., Home, SNF, Rehab)",
  "admission_diagnosis": "Primary diagnosis at time of admission",
  "discharge_diagnosis": "Final diagnosis at time of discharge",
  "secondary_diagnoses": ["List of secondary diagnoses"],
  "presenting_complaints": "Chief complaints on admission",
  "hospital_course": "Narrative summary of the hospital course and treatment",
  "procedures_performed": ["List of procedures performed during admission"],
  "significant_results": ["Key lab/imaging results that guided management"],
  "discharge_medications": [
    {"medication": "name", "dose": "dose", "frequency": "frequency", "duration": "duration", "notes": "notes"}
  ],
  "medications_stopped": ["List of medications that were stopped and reason"],
  "allergies": ["List of known allergies"],
  "discharge_condition": "Patient condition at discharge (Stable/Improved/etc.)",
  "discharge_instructions": ["List of instructions given to patient"],
  "activity_restrictions": "Activity level and restrictions post-discharge",
  "diet_instructions": "Dietary instructions post-discharge",
  "wound_care": "Wound care instructions if applicable",
  "follow_up_appointments": [
    {"provider": "provider name", "specialty": "specialty", "timeframe": "within X days", "reason": "reason"}
  ],
  "return_to_ed_criteria": ["List of symptoms/signs that should prompt ED return"],
  "patient_education": ["Topics discussed with patient/family prior to discharge"],
  "code_status": "Code status during admission",
  "additional_notes": "Any other relevant information"
}

Extract as much information as possible from the provided notes. Use empty strings or empty arrays for fields that cannot be determined. Return only valid JSON."""


def generate_discharge_summary(
    admission_note: str,
    treatment_course: str = "",
    discharge_plan: str = "",
) -> dict:
    """
    Generate a comprehensive discharge summary document.

    Args:
        admission_note: Admission notes and presenting complaints
        treatment_course: Summary of treatment provided during hospitalization
        discharge_plan: Discharge plan including medications and follow-up

    Returns:
        dict with structured discharge summary data and a generated DOCX file
    """
    client = OpenAI(api_key=os.environ.get("OPENAI_API_KEY"))
    model = os.environ.get("OPENAI_MODEL", "gpt-4o")

    parts = [f"## Admission Note\n{admission_note}"]
    if treatment_course:
        parts.append(f"## Treatment Course\n{treatment_course}")
    if discharge_plan:
        parts.append(f"## Discharge Plan\n{discharge_plan}")
    combined_input = "\n\n".join(parts)

    logger.info("Generating discharge summary via OpenAI")

    try:
        response = client.chat.completions.create(
            model=model,
            temperature=0.0,
            messages=[
                {"role": "system", "content": DISCHARGE_SYSTEM_PROMPT},
                {"role": "user", "content": combined_input},
            ],
            response_format={"type": "json_object"},
        )

        result_text = response.choices[0].message.content
        summary_data = json.loads(result_text)

        logger.info("Discharge summary generated successfully")

        docx_b64, filename = _create_discharge_docx(summary_data)

        return {
            "summary": summary_data,
            "filename": filename,
            "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "content_base64": docx_b64,
            "message": f"Discharge summary generated successfully: {filename}",
        }

    except Exception as e:
        logger.error(f"Error generating discharge summary: {e}")
        raise


def _create_discharge_docx(summary_data: dict) -> tuple[str, str]:
    """Create a formatted DOCX discharge summary from structured data."""
    import io
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Title
    title = doc.add_heading("Discharge Summary", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Patient header
    doc.add_paragraph(
        f"Patient: {summary_data.get('patient_name', 'N/A')}  |  "
        f"MRN: {summary_data.get('patient_id', 'N/A')}  |  "
        f"DOB: {summary_data.get('patient_dob', 'N/A')}"
    )
    doc.add_paragraph(
        f"Admission: {summary_data.get('admission_date', 'N/A')}  |  "
        f"Discharge: {summary_data.get('discharge_date', 'N/A')}  |  "
        f"LOS: {summary_data.get('length_of_stay', 'N/A')} days"
    )
    doc.add_paragraph(
        f"Attending: {summary_data.get('attending_physician', 'N/A')}  |  "
        f"Ward: {summary_data.get('ward', 'N/A')}  |  "
        f"Disposition: {summary_data.get('discharge_disposition', 'N/A')}"
    )
    doc.add_paragraph()

    def add_section(heading: str, content):
        doc.add_heading(heading, level=2)
        if isinstance(content, list):
            if content:
                for item in content:
                    if isinstance(item, dict):
                        # Format dict items (e.g., medications)
                        text = ", ".join(f"{k}: {v}" for k, v in item.items() if v)
                        doc.add_paragraph(f"• {text}", style="List Bullet")
                    else:
                        doc.add_paragraph(f"• {item}", style="List Bullet")
            else:
                doc.add_paragraph("None documented.")
        else:
            doc.add_paragraph(str(content) if content else "Not documented.")

    add_section("Admission Diagnosis", summary_data.get("admission_diagnosis", ""))
    add_section("Discharge Diagnosis", summary_data.get("discharge_diagnosis", ""))
    add_section("Secondary Diagnoses", summary_data.get("secondary_diagnoses", []))
    add_section("Presenting Complaints", summary_data.get("presenting_complaints", ""))
    add_section("Hospital Course", summary_data.get("hospital_course", ""))
    add_section("Procedures Performed", summary_data.get("procedures_performed", []))
    add_section("Significant Results", summary_data.get("significant_results", []))
    add_section("Allergies", summary_data.get("allergies", []))

    # Discharge medications (special formatting)
    doc.add_heading("Discharge Medications", level=2)
    meds = summary_data.get("discharge_medications", [])
    if meds:
        for med in meds:
            if isinstance(med, dict):
                name = med.get("medication", "")
                dose = med.get("dose", "")
                freq = med.get("frequency", "")
                dur = med.get("duration", "")
                notes = med.get("notes", "")
                parts = [name, dose, freq]
                if dur:
                    parts.append(f"for {dur}")
                line = " — ".join(filter(None, parts))
                if notes:
                    line += f" ({notes})"
                doc.add_paragraph(f"• {line}", style="List Bullet")
    else:
        doc.add_paragraph("No medications at discharge.")

    add_section("Medications Stopped", summary_data.get("medications_stopped", []))
    add_section("Discharge Condition", summary_data.get("discharge_condition", ""))
    add_section("Discharge Instructions", summary_data.get("discharge_instructions", []))
    add_section("Activity Restrictions", summary_data.get("activity_restrictions", ""))
    add_section("Diet Instructions", summary_data.get("diet_instructions", ""))
    add_section("Wound Care", summary_data.get("wound_care", ""))

    # Follow-up appointments
    doc.add_heading("Follow-up Appointments", level=2)
    followups = summary_data.get("follow_up_appointments", [])
    if followups:
        for appt in followups:
            if isinstance(appt, dict):
                provider = appt.get("provider", "")
                specialty = appt.get("specialty", "")
                timeframe = appt.get("timeframe", "")
                reason = appt.get("reason", "")
                line = f"{provider} ({specialty}) — {timeframe}"
                if reason:
                    line += f": {reason}"
                doc.add_paragraph(f"• {line}", style="List Bullet")
    else:
        doc.add_paragraph("No follow-up appointments specified.")

    add_section("Return to ED If", summary_data.get("return_to_ed_criteria", []))
    add_section("Patient Education", summary_data.get("patient_education", []))
    add_section("Code Status", summary_data.get("code_status", ""))
    add_section("Additional Notes", summary_data.get("additional_notes", ""))

    buf = io.BytesIO()
    doc.save(buf)
    doc_bytes = buf.getvalue()

    patient_name = summary_data.get("patient_name", "patient").replace(" ", "_")
    date = summary_data.get("discharge_date", "").replace("-", "") or "undated"
    filename = f"discharge_summary_{patient_name}_{date}.docx"

    return base64.b64encode(doc_bytes).decode("utf-8"), filename